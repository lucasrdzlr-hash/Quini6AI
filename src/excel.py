from openpyxl import load_workbook
from .models import Numero

def leer_stats(p):
 wb=load_workbook(p,data_only=True);ws=wb['Stats'];h={c.value:i+1 for i,c in enumerate(ws[1])};o=[]
 for r in range(2,ws.max_row+1):
  if ws.cell(r,h['numero']).value is None: continue
  o.append(Numero(ws.cell(r,h['numero']).value,ws.cell(r,h['ap20']).value,ws.cell(r,h['au20']).value,ws.cell(r,h['sc20']).value,ws.cell(r,h['ap50']).value,ws.cell(r,h['au50']).value,ws.cell(r,h['sc50']).value,ws.cell(r,h['ap100']).value,ws.cell(r,h['au100']).value,ws.cell(r,h['sc100']).value,ws.cell(r,h['hist_ap']).value,str(ws.cell(r,h['ult_salida']).value),ws.cell(r,h['atraso']).value,str(ws.cell(r,h['grupo']).value),str(ws.cell(r,h['categoria']).value)))
 return o
